#!/usr/bin/env python3
"""Extract text + metadata from knowledge files (PDF, PPTX)."""

import argparse
import json
import re
import subprocess
import sys
from datetime import date
from pathlib import Path

from pptx import Presentation

SUPPORTED = {".pdf", ".pptx"}

MONTHS = {
    "jan": 1, "january": 1, "feb": 2, "february": 2,
    "mar": 3, "march": 3, "apr": 4, "april": 4,
    "may": 5, "jun": 6, "june": 6,
    "jul": 7, "july": 7, "aug": 8, "august": 8,
    "sep": 9, "sept": 9, "september": 9,
    "oct": 10, "october": 10, "nov": 11, "november": 11,
    "dec": 12, "december": 12,
}
_MONTH_RE = "|".join(sorted(MONTHS, key=len, reverse=True))


# ── date extraction ──────────────────────────────────────────────

def extract_date(path: str) -> dict:
    years = [int(y) for y in re.findall(r"\b(20[12]\d)\b", path)]

    # Month DD [, YYYY] — "Oct 28" or "January 15, 2024"
    for m in re.finditer(rf"\b({_MONTH_RE})\.?\s+(\d{{1,2}})\b", path, re.I):
        day = int(m.group(2))
        if not 1 <= day <= 31:
            continue
        ym = re.match(r"[,\s]+(20[12]\d)\b", path[m.end():])
        year = int(ym.group(1)) if ym else (years[-1] if years else None)
        if year:
            try:
                return {"date": date(year, MONTHS[m.group(1).lower()], day).isoformat(), "precision": "day"}
            except ValueError:
                pass

    # Month YYYY — "March 2026"
    m = re.search(rf"\b({_MONTH_RE})\.?\s+(20[12]\d)\b", path, re.I)
    if m:
        return {"date": date(int(m.group(2)), MONTHS[m.group(1).lower()], 1).isoformat(), "precision": "month"}

    # YYYYMM_ prefix — "202604_"
    m = re.search(r"\b(20[12]\d)(0[1-9]|1[0-2])_", path)
    if m:
        return {"date": date(int(m.group(1)), int(m.group(2)), 1).isoformat(), "precision": "month"}

    # QN [YYYY] — "Q1 2025" or bare "Q1" with year from path
    m = re.search(r"\bQ([1-4])(?:\s+(20[12]\d))?\b", path, re.I)
    if m:
        year = int(m.group(2)) if m.group(2) else (years[-1] if years else None)
        if year:
            return {"date": date(year, (int(m.group(1)) - 1) * 3 + 1, 1).isoformat(), "precision": "quarter"}

    # HN [YYYY] — "H2 2025" or bare "H1" with year from path
    m = re.search(r"\bH([12])(?:\s+(20[12]\d))?\b", path)
    if m:
        year = int(m.group(2)) if m.group(2) else (years[-1] if years else None)
        if year:
            return {"date": date(year, 1 if m.group(1) == "1" else 7, 1).isoformat(), "precision": "half"}

    # Bare year
    if years:
        return {"date": date(years[-1], 1, 1).isoformat(), "precision": "year"}

    return {"date": None, "precision": None}


# ── text extraction ──────────────────────────────────────────────

def extract_pdf(src: Path, dst: Path) -> bool:
    r = subprocess.run(["pdftotext", str(src), str(dst)], capture_output=True, text=True)
    if r.returncode != 0:
        print(f"FAIL: {r.stderr.strip()}", file=sys.stderr)
        return False
    return True


def _slide_text(slide) -> list[str]:
    lines = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    lines.append(text)
        if shape.has_table:
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
    return lines


def extract_pptx(src: Path, dst: Path) -> bool:
    try:
        prs = Presentation(str(src))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            text = _slide_text(slide)
            if text:
                parts.append(f"[Slide {i}]")
                parts.extend(text)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append("[Notes]")
                    parts.append(notes)
            parts.append("")
        dst.write_text("\n".join(parts).rstrip() + "\n")
        return True
    except Exception as e:
        print(f"FAIL: {e}", file=sys.stderr)
        return False


def extract_pptx_metadata(src: Path) -> dict:
    try:
        prs = Presentation(str(src))
        props = prs.core_properties
        result = {}
        if props.author:
            result["author"] = props.author
        if props.title:
            result["title"] = props.title
        if props.created:
            result["created"] = props.created.isoformat()
        if props.modified:
            result["modified"] = props.modified.isoformat()
        return result
    except Exception:
        return {}


EXTRACTORS = {".pdf": extract_pdf, ".pptx": extract_pptx}


# ── slugify ──────────────────────────────────────────────────────

def slugify(path: str) -> str:
    name = path
    for s in SUPPORTED:
        if name.lower().endswith(s):
            name = name[: -len(s)]
            break
    name = name.lower().replace("/", "-").replace("\\", "-")
    name = re.sub(r"[^a-z0-9\-]", "-", name)
    name = re.sub(r"-{2,}", "-", name).strip("-")
    return name + ".txt"


# ── main ─────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(description="Extract text + metadata from knowledge files.")
    ap.add_argument("src", help="Source directory")
    ap.add_argument("dst", help="Output directory")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    src = Path(args.src).resolve()
    dst = Path(args.dst).resolve()
    if not src.is_dir():
        sys.exit(f"Not a directory: {src}")

    files = sorted(f for f in src.rglob("*") if f.suffix.lower() in SUPPORTED and f.is_file())
    if not files:
        sys.exit(f"No supported files in {src}")

    dst.mkdir(parents=True, exist_ok=True)
    seen: dict[str, str] = {}
    manifest: list[dict] = []
    ok = fail = 0

    for f in files:
        rel = str(f.relative_to(src))
        slug = slugify(rel)
        if slug in seen:
            print(f"  SKIP (collision): {rel}  ->  {slug}  (vs {seen[slug]})", file=sys.stderr)
            fail += 1
            continue
        seen[slug] = rel

        date_info = extract_date(rel)
        entry = {
            "source_path": rel,
            "output_file": slug,
            "file_type": f.suffix.lower().lstrip("."),
            "date": date_info["date"],
            "date_precision": date_info["precision"],
        }
        if f.suffix.lower() == ".pptx":
            embedded = extract_pptx_metadata(f)
            if embedded:
                entry["embedded"] = embedded
        manifest.append(entry)

        if args.dry_run:
            d = date_info["date"] or "no date"
            p = date_info["precision"] or "-"
            print(f"  {rel}\n    -> {slug}  [{d} ({p})]")
            continue

        print(f"  {rel} -> {slug}", end=" ... ")
        if EXTRACTORS[f.suffix.lower()](f, dst / slug):
            print("OK")
            ok += 1
        else:
            fail += 1

    (dst / "metadata.json").write_text(json.dumps(manifest, indent=2) + "\n")

    if args.dry_run:
        print(f"\n{len(manifest)} files planned")
    else:
        print(f"\nDone: {ok} ok, {fail} failed, {len(manifest)} total")
    print(f"Manifest: {dst / 'metadata.json'}")


if __name__ == "__main__":
    main()

from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from knowledge_extract import extract_date, slugify

SCRIPT = str(Path(__file__).resolve().parent.parent / "knowledge_extract.py")


# ── helpers ──────────────────────────────────────────────────────

def make_pptx(path: Path, slides: list[str], notes: list[str] | None = None):
    prs = Presentation()
    for i, text in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
        txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
        txBox.text_frame.text = text
        if notes and i < len(notes) and notes[i]:
            slide.notes_slide.notes_text_frame.text = notes[i]
    prs.save(str(path))


def make_pptx_with_table(path: Path, rows: list[list[str]]):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    n_rows, n_cols = len(rows), len(rows[0])
    table = slide.shapes.add_table(n_rows, n_cols, Inches(1), Inches(1), Inches(6), Inches(2)).table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    prs.save(str(path))


def make_pdf(path: Path, text: str = "Hello World"):
    content = f"BT /F1 12 Tf 100 700 Td ({text}) Tj ET"
    obj = [
        "1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj",
        "2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj",
        (
            "3 0 obj\n<< /Type /Page /Parent 2 0 R "
            "/MediaBox [0 0 612 792] /Contents 4 0 R "
            "/Resources << /Font << /F1 5 0 R >> >> >>\nendobj"
        ),
        f"4 0 obj\n<< /Length {len(content)} >>\nstream\n{content}\nendstream\nendobj",
        "5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj",
    ]
    body = "%PDF-1.0\n"
    offsets = []
    for o in obj:
        offsets.append(len(body))
        body += o + "\n"
    xref_pos = len(body)
    body += f"xref\n0 {len(obj) + 1}\n0000000000 65535 f \n"
    for off in offsets:
        body += f"{off:010d} 00000 n \n"
    body += f"trailer\n<< /Size {len(obj) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n"
    path.write_text(body)


HAS_PDFTOTEXT = shutil.which("pdftotext") is not None


# ── date extraction ──────────────────────────────────────────────

class TestDateExtraction:
    def test_month_day_year_from_path(self):
        r = extract_date("All Hands/2025/H2 2025/11. Oct 28 All-Hands.pptx")
        assert r == {"date": "2025-10-28", "precision": "day"}

    def test_month_day_year_inline(self):
        r = extract_date("Report - January 15, 2024.pdf")
        assert r == {"date": "2024-01-15", "precision": "day"}

    def test_month_year(self):
        r = extract_date("Commercial Strategy Memo March 2026.pdf")
        assert r == {"date": "2026-03-01", "precision": "month"}

    def test_yyyymm_prefix(self):
        r = extract_date("202604_Sona launch strategy.pdf")
        assert r == {"date": "2026-04-01", "precision": "month"}

    def test_quarter_with_year(self):
        r = extract_date("VKO decks/2025/Q1 Updates.pptx")
        assert r == {"date": "2025-01-01", "precision": "quarter"}

    def test_quarter_year_from_path(self):
        r = extract_date("decks/2025/Q3 Review.pptx")
        assert r == {"date": "2025-07-01", "precision": "quarter"}

    def test_half_with_year(self):
        r = extract_date("All Hands/H2 2025/deck.pptx")
        assert r == {"date": "2025-07-01", "precision": "half"}

    def test_half_year_from_path(self):
        r = extract_date("All Hands/2026/H1/deck.pptx")
        assert r == {"date": "2026-01-01", "precision": "half"}

    def test_bare_year(self):
        r = extract_date("VKO decks/2023/Kickoff.pdf")
        assert r == {"date": "2023-01-01", "precision": "year"}

    def test_no_date(self):
        r = extract_date("Handbook/Compensation.pdf")
        assert r == {"date": None, "precision": None}

    def test_feb_short(self):
        r = extract_date("2023/H1 2023/2. Feb 13 All-Hands Meeting/deck.pdf")
        assert r == {"date": "2023-02-13", "precision": "day"}

    def test_september_long(self):
        r = extract_date("2021/3. September 20 All-Hands/deck.pptx")
        assert r == {"date": "2021-09-20", "precision": "day"}


# ── slugify ──────────────────────────────────────────────────────

class TestSlugify:
    def test_basic(self):
        assert slugify("Report.pdf") == "report.txt"

    def test_nested_path(self):
        assert slugify("A/B/C.pptx") == "a-b-c.txt"

    def test_special_chars(self):
        assert slugify("G&A Strategy Memo.pdf") == "g-a-strategy-memo.txt"

    def test_emoji(self):
        s = slugify("📗 US Employee Handbook/Overview.pdf")
        assert s == "us-employee-handbook-overview.txt"

    def test_no_leading_trailing_hyphens(self):
        s = slugify("  spaced out .pdf")
        assert not s.startswith("-")
        assert s.endswith(".txt")


# ── pptx extraction ──────────────────────────────────────────────

class TestPptxExtraction:
    def test_basic_text(self, tmp_path):
        src = tmp_path / "src"
        src.mkdir()
        make_pptx(src / "deck.pptx", ["Slide one content", "Slide two content"])
        dst = tmp_path / "dst"

        subprocess.run([sys.executable, SCRIPT, str(src), str(dst)], check=True)

        out = (dst / "deck.txt").read_text()
        assert "Slide one content" in out
        assert "Slide two content" in out
        assert "[Slide 1]" in out
        assert "[Slide 2]" in out

    def test_speaker_notes(self, tmp_path):
        src = tmp_path / "src"
        src.mkdir()
        make_pptx(
            src / "deck.pptx",
            ["Visible text"],
            notes=["These are the speaker notes"],
        )
        dst = tmp_path / "dst"

        subprocess.run([sys.executable, SCRIPT, str(src), str(dst)], check=True)

        out = (dst / "deck.txt").read_text()
        assert "[Notes]" in out
        assert "speaker notes" in out

    def test_table_content(self, tmp_path):
        src = tmp_path / "src"
        src.mkdir()
        make_pptx_with_table(src / "deck.pptx", [["Name", "Score"], ["Alice", "95"]])
        dst = tmp_path / "dst"

        subprocess.run([sys.executable, SCRIPT, str(src), str(dst)], check=True)

        out = (dst / "deck.txt").read_text()
        assert "Alice" in out
        assert "95" in out


# ── pdf extraction ───────────────────────────────────────────────

@pytest.mark.skipif(not HAS_PDFTOTEXT, reason="pdftotext not installed")
class TestPdfExtraction:
    def test_basic_text(self, tmp_path):
        src = tmp_path / "src"
        src.mkdir()
        make_pdf(src / "doc.pdf", "Test document content")
        dst = tmp_path / "dst"

        subprocess.run([sys.executable, SCRIPT, str(src), str(dst)], check=True)

        out = (dst / "doc.txt").read_text()
        assert "Test document content" in out


# ── full pipeline ────────────────────────────────────────────────

class TestFullPipeline:
    def test_mixed_directory(self, tmp_path):
        src = tmp_path / "src"
        (src / "2025" / "H2 2025").mkdir(parents=True)
        (src / "handbook").mkdir()

        make_pptx(
            src / "2025" / "H2 2025" / "Oct 28 All-Hands.pptx",
            ["All hands slide"],
        )
        make_pptx(src / "handbook" / "Onboarding.pptx", ["Welcome"])

        dst = tmp_path / "dst"
        subprocess.run([sys.executable, SCRIPT, str(src), str(dst)], check=True)

        manifest = json.loads((dst / "metadata.json").read_text())
        assert len(manifest) == 2

        by_src = {e["source_path"]: e for e in manifest}

        allhands = by_src["2025/H2 2025/Oct 28 All-Hands.pptx"]
        assert allhands["date"] == "2025-10-28"
        assert allhands["date_precision"] == "day"
        assert allhands["file_type"] == "pptx"

        onboarding = by_src["handbook/Onboarding.pptx"]
        assert onboarding["date"] is None

        assert (dst / allhands["output_file"]).exists()
        assert (dst / onboarding["output_file"]).exists()

    def test_dry_run_creates_no_text_files(self, tmp_path):
        src = tmp_path / "src"
        src.mkdir()
        make_pptx(src / "deck.pptx", ["content"])
        dst = tmp_path / "dst"

        subprocess.run(
            [sys.executable, SCRIPT, str(src), str(dst), "--dry-run"],
            check=True,
        )

        assert (dst / "metadata.json").exists()
        txt_files = list(dst.glob("*.txt"))
        assert txt_files == []

    def test_slug_collision_skips(self, tmp_path):
        src = tmp_path / "src"
        # "a/deck.pptx" and "a-deck.pptx" both slugify to "a-deck.txt"
        (src / "a").mkdir(parents=True)
        make_pptx(src / "a" / "deck.pptx", ["first"])
        make_pptx(src / "a-deck.pptx", ["second"])
        dst = tmp_path / "dst"

        result = subprocess.run(
            [sys.executable, SCRIPT, str(src), str(dst)],
            capture_output=True, text=True,
        )

        assert "SKIP (collision)" in result.stderr
        manifest = json.loads((dst / "metadata.json").read_text())
        assert len(manifest) == 1

    @pytest.mark.skipif(not HAS_PDFTOTEXT, reason="pdftotext not installed")
    def test_mixed_pdf_and_pptx(self, tmp_path):
        src = tmp_path / "src" / "Monthly Memos"
        src.mkdir(parents=True)
        make_pdf(src / "Commercial Strategy Memo March 2026.pdf", "Revenue grew")
        make_pptx(src / "April 2026 Review.pptx", ["Q1 results"])
        dst = tmp_path / "dst"

        subprocess.run([sys.executable, SCRIPT, str(tmp_path / "src"), str(dst)], check=True)

        manifest = json.loads((dst / "metadata.json").read_text())
        by_src = {e["source_path"]: e for e in manifest}

        pdf_entry = by_src["Monthly Memos/Commercial Strategy Memo March 2026.pdf"]
        assert pdf_entry["date"] == "2026-03-01"
        assert pdf_entry["date_precision"] == "month"

        pptx_entry = by_src["Monthly Memos/April 2026 Review.pptx"]
        assert pptx_entry["date"] == "2026-04-01"
        assert pptx_entry["date_precision"] == "month"

        pdf_txt = (dst / pdf_entry["output_file"]).read_text()
        assert "Revenue grew" in pdf_txt
        pptx_txt = (dst / pptx_entry["output_file"]).read_text()
        assert "Q1 results" in pptx_txt
